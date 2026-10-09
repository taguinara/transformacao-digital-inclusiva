# -*- coding: utf-8 -*-
"""
Converte o Controle_Vocabulario.docx em Excel tabular,
pronto para o pipeline de curadoria.

Trata:
  - múltiplas linhas de cabeçalho
  - células mescladas (herança de Área / Subárea)
  - quebras de linha internas
  - remoção de linhas vazias
"""

from pathlib import Path
import pandas as pd
from docx import Document


DOCX = "controle_vocabulario.docx"
XLSX = "controle_vocabulario.xlsx"

# Nomes das colunas de saída (na ordem em que aparecem no Word)
COLUNAS_SAIDA = [
    "Area", "Subarea", "Tema_Escopo",
    "Area_Termo", "Subarea_Termo",
    "Garantia_Literaria", "Remissivas",
]

# Palavras-chave que identificam a LINHA DE CABEÇALHO real
# (o Word tem 1 ou 2 linhas de cabeçalho antes dos dados)
CHAVES_HEADER = ["área", "subárea", "garantia"]


# ----------------------------------------------------------------------
# 1. Localizar a linha de cabeçalho
# ----------------------------------------------------------------------
def encontrar_cabecalho(tabela) -> int:
    """Retorna o índice da linha que contém os nomes das colunas."""
    for i, linha in enumerate(tabela.rows[:5]):  # procura só no topo
        textos = " ".join(c.text.lower() for c in linha.cells)
        if all(k in textos for k in CHAVES_HEADER):
            return i
    raise ValueError(
        "Não foi possível localizar a linha de cabeçalho na tabela. "
        "Confira se o .docx tem as colunas Área/Subárea/Garantia literária."
    )


# ----------------------------------------------------------------------
# 2. Extrair dados da tabela
# ----------------------------------------------------------------------
def extrair_linhas(tabela, idx_cabecalho: int) -> list[list[str]]:
    """
    Percorre as linhas após o cabeçalho.
    Células mescladas aparecem repetidas — deduplicamos pelo id do XML.
    """
    dados = []
    for linha in tabela.rows[idx_cabecalho + 1:]:
        celulas = []
        vistos = set()
        for celula in linha.cells:
            # id() do nó XML identifica células mescladas
            chave = id(celula._tc)
            if chave in vistos:
                continue
            vistos.add(chave)
            # Junta parágrafos com espaço, remove quebras internas
            texto = " ".join(p.text.strip() for p in celula.paragraphs)
            texto = " ".join(texto.split())  # colapsa espaços
            celulas.append(texto)
        dados.append(celulas)
    return dados


# ----------------------------------------------------------------------
# 3. Alinhar larguras (o Word pode devolver menos colunas por causa
#    de mesclagens horizontais). Completamos com "" à direita.
# ----------------------------------------------------------------------
def alinhar(dados: list[list[str]], n_colunas: int) -> list[list[str]]:
    return [linha + [""] * (n_colunas - len(linha)) for linha in dados]


# ----------------------------------------------------------------------
# 4. Herdar Área e Subárea das linhas anteriores
#    (no Word, elas só aparecem na primeira linha do bloco)
# ----------------------------------------------------------------------
def preencher_hierarquia(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["Area"] = df["Area"].replace("", pd.NA).ffill().fillna("")
    df["Subarea"] = df["Subarea"].replace("", pd.NA).ffill().fillna("")
    return df


# ----------------------------------------------------------------------
# 5. Pipeline completo
# ----------------------------------------------------------------------
def converter(docx_path=DOCX, xlsx_path=XLSX) -> pd.DataFrame:
    docx_path = Path(docx_path)
    if not docx_path.exists():
        raise FileNotFoundError(f"'{docx_path}' não encontrado.")

    doc = Document(docx_path)
    if not doc.tables:
        raise ValueError("O documento não contém nenhuma tabela.")

    tabela = doc.tables[0]
    print(f"✔ Tabela encontrada:{len(tabela.rows)} linhas brutas")

    idx_cab = encontrar_cabecalho(tabela)
    print(f"✔ Linha de cabeçalho detectada: índice{idx_cab}")

    linhas = extrair_linhas(tabela, idx_cab)
    linhas = alinhar(linhas, len(COLUNAS_SAIDA))

    df = pd.DataFrame(linhas, columns=COLUNAS_SAIDA)
    df = preencher_hierarquia(df)

    # Remove linhas totalmente vazias (rodapés, espaçadores)
    antes = len(df)
    df = df[df.apply(lambda r: any(v.strip() for v in r), axis=1)]
    df = df.drop_duplicates().reset_index(drop=True)
    print(f"✔{antes} →{len(df)} linhas após limpeza")

    df.to_excel(xlsx_path, index=False)
    print(f"✔ Excel gerado:{xlsx_path}")
    return df


# ----------------------------------------------------------------------
# 6. Execução
# ----------------------------------------------------------------------
if __name__ == "__main__":
    df = converter()

    print("\n--- Amostra (5 primeiras linhas) ---")
    print(df.head().to_string(index=False))

    print("\n--- Resumo ---")
    print(f"Total de registros:{len(df)}")
    print("\nLinhas por Área:")
    print(df["Area"].value_counts().to_string())